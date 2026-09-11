# -*- coding: utf-8 -*-
"""文章多格式转换: html -> pdf / txt / md / word(docx)

设计:
  - 四种格式全部从已保存的 html 转换(html 是中间产物)
  - html 图片目录: images/(现状, localize_article_images 生成)
  - md 图片目录: mdimgs/(md 专属外挂图片夹, 从 images/ 复制)
  - 若 formats 未选择 html: 全部转换完成后删除 html + images/(md 选了就保留 mdimgs)
"""
import os
import re
import shutil
import subprocess
import pathlib
from html.parser import HTMLParser

from ..core import logkit

log = logkit.get_logger("collect.formats")   # 多格式转换日志 -> run.log

_EDGE_CANDS = [
    r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
]


def find_edge():
    """定位系统 Edge(打印 PDF 用)"""
    for p in _EDGE_CANDS:
        if os.path.isfile(p):
            return p
    return None


# ---------------- html 文本提取(正文优先 js_content) ----------------

class _TextExtract(HTMLParser):
    """提取文本: 优先 <div id="js_content"> 正文容器; 找不到则全文档(去脚本/样式)"""
    SKIP = {"script", "style", "noscript"}
    BLOCK = {"p", "br", "div", "li", "h1", "h2", "h3", "h4", "section", "tr"}

    def __init__(self, body_only=True):
        super().__init__(convert_charrefs=True)
        self.body_only = body_only
        self.in_body = not body_only
        self.depth = 0          # js_content 内嵌套深度
        self.skip = 0
        self.parts = []
        self.imgs = []          # 正文图片 src 列表
        self.title = ""

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag in self.SKIP:
            self.skip += 1
            return
        if not self.in_body:
            if tag == "div" and a.get("id") == "js_content":
                self.in_body = True
                self.depth = 1
                return
            # 标题兜底
            if tag == "title":
                self._grab_title = True
            return
        if tag == "div":
            self.depth += 1
        if tag in self.BLOCK:
            self.parts.append("\n")
        if tag == "img":
            src = a.get("src") or a.get("data-src") or ""
            if src and (src.startswith("images/") or src.lower().endswith((".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp"))):
                self.imgs.append(src)

    def handle_endtag(self, tag):
        if tag in self.SKIP and self.skip > 0:
            self.skip -= 1
            return
        if not self.in_body:
            return
        if tag == "div":
            self.depth -= 1
            if self.depth <= 0:
                self.in_body = False

    def handle_data(self, d):
        if self.skip:
            return
        if getattr(self, "_grab_title", False) and d.strip():
            self.title = d.strip()
        if self.in_body and d.strip():
            self.parts.append(d)

    def text(self):
        return "\n".join(l.strip() for l in "".join(self.parts).splitlines() if l.strip())


def extract_text(html):
    """html -> 纯文本(正文优先)"""
    p = _TextExtract(body_only=True)
    try:
        p.feed(html)
    except Exception:
        pass
    txt = p.text()
    if len(txt) < 50:                      # 正文容器没命中 -> 退化为全文档
        p2 = _TextExtract(body_only=False)
        try:
            p2.feed(html)
        except Exception:
            pass
        txt = p2.text()
    return txt, p.imgs


def _meta(html, html_path):
    """标题/公众号/日期: 优先 html meta/目录名(日期_标题)"""
    title = ""
    m = re.search(r'<meta[^>]+property=["\']og:title["\'][^>]+content=["\']([^"\']+)', html, re.I)
    if m:
        title = m.group(1).strip()
    if not title:
        m = re.search(r"<title>([^<]+)</title>", html, re.I)
        if m:
            title = re.sub(r"\s+", " ", m.group(1)).strip()
    folder = os.path.basename(os.path.dirname(html_path))       # <日期>_<标题>
    acc = os.path.basename(os.path.dirname(os.path.dirname(html_path)))
    date = folder.split("_", 1)[0] if "_" in folder else ""
    if not title:
        title = folder.split("_", 1)[-1]
    return title, acc, date


# ---------------- 四种格式转换 ----------------

def html_to_txt(html_path, txt_path=None, html=None):
    """html -> txt(纯文本, 正文优先)"""
    html = html if html is not None else open(html_path, encoding="utf-8", errors="ignore").read()
    txt, _ = extract_text(html)
    txt_path = txt_path or os.path.splitext(html_path)[0] + ".txt"
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(txt)
    return txt_path


def html_to_md(html_path, md_path=None, html=None, keep_imgs=True):
    """html -> md(标题/元信息/正文; 图片复制到 mdimgs/ 并引用)"""
    html = html if html is not None else open(html_path, encoding="utf-8", errors="ignore").read()
    title, acc, date = _meta(html, html_path)
    txt, imgs = extract_text(html)
    art_dir = os.path.dirname(html_path)
    md_path = md_path or os.path.splitext(html_path)[0] + ".md"
    lines = [f"# {title}", ""]
    meta = []
    if acc:
        meta.append(f"公众号: {acc}")
    if date:
        meta.append(f"日期: {date}")
    if meta:
        lines.append("> " + " | ".join(meta))
        lines.append("")
    # 图片复制到 mdimgs/(去重保序); 源目录兼容 images/ 与 <stem>_files/
    dst_dir = os.path.join(art_dir, "mdimgs")
    seen = []
    for s in imgs:
        fn = os.path.basename(s)
        if fn and fn not in seen:
            seen.append(fn)
    img_srcs = []
    for s in imgs:
        p_f = os.path.join(art_dir, s)
        if os.path.isfile(p_f):
            img_srcs.append(p_f)
    img_srcs = list(dict.fromkeys(img_srcs))   # 保序去重
    if img_srcs:
        if keep_imgs:
            os.makedirs(dst_dir, exist_ok=True)
        for p_src in img_srcs:
            fn = os.path.basename(p_src)
            if keep_imgs:
                try:
                    shutil.copy2(p_src, os.path.join(dst_dir, fn))
                except Exception as e:
                    log.warning("[md] 图片复制失败 %s: %s", fn, e)
            lines.append(f"![{fn}](mdimgs/{fn})")
            lines.append("")
    lines.append(txt)
    with open(md_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    return md_path


def html_to_pdf(html_path, pdf_path=None, edge=None, wait_ms=8000):
    """html -> pdf(Edge headless --print-to-pdf; virtual-time-budget 等资源加载)"""
    edge = edge or find_edge()
    if not edge:
        raise RuntimeError("未找到 Edge 浏览器, 无法转 PDF")
    pdf_path = pdf_path or os.path.splitext(html_path)[0] + ".pdf"
    url = pathlib.Path(html_path).as_uri()
    r = subprocess.run(
        [edge, "--headless", "--disable-gpu",
         f"--virtual-time-budget={wait_ms}",
         "--print-to-pdf=" + pdf_path, url],
        capture_output=True, timeout=180)
    if r.returncode != 0 or not os.path.exists(pdf_path) or os.path.getsize(pdf_path) == 0:
        raise RuntimeError(f"Edge 转 PDF 失败 rc={r.returncode}")
    return pdf_path


def html_to_docx(html_path, docx_path=None, html=None):
    """html -> word(docx): 标题 + 元信息 + 正文段落(图片按序嵌入, 失败跳过)"""
    from docx import Document
    from docx.shared import Inches
    html = html if html is not None else open(html_path, encoding="utf-8", errors="ignore").read()
    title, acc, date = _meta(html, html_path)
    txt, imgs = extract_text(html)
    docx_path = docx_path or os.path.splitext(html_path)[0] + ".docx"
    doc = Document()
    doc.add_heading(title or "未命名文章", level=1)
    meta = " | ".join(x for x in [f"公众号: {acc}" if acc else "", f"日期: {date}" if date else ""] if x)
    if meta:
        doc.add_paragraph(meta)
    for line in txt.splitlines():
        line = line.strip()
        if line:
            doc.add_paragraph(line)
    # 图片嵌入(正文末尾, 失败跳过)
    art_dir = os.path.dirname(html_path)
    for s in imgs:
        src = os.path.join(art_dir, s)
        if os.path.isfile(src):
            try:
                doc.add_picture(src, width=Inches(5.5))
            except Exception as e:
                log.warning("[word] 图片嵌入失败 %s: %s", s, e)
    doc.save(docx_path)
    return docx_path


# ---------------- 统一入口 ----------------

_FMT_FUNCS = {
    "txt": html_to_txt,
    "md": html_to_md,
    "pdf": html_to_pdf,
    "word": html_to_docx,
    "docx": html_to_docx,
}


def convert_html(html_path, formats):
    """按 formats 转换 html -> 各格式文件; 返回 {fmt: path}(失败跳过并记日志)"""
    out = {}
    html = open(html_path, encoding="utf-8", errors="ignore").read()   # 一次读取复用
    for fmt in formats:
        fmt = (fmt or "").strip().lower()
        if fmt in ("html", ""):
            continue
        fn = _FMT_FUNCS.get(fmt)
        if not fn:
            log.warning("[formats] 未知格式: %s", fmt)
            continue
        try:
            if fmt in ("md", "pdf"):
                path = fn(html_path, html=html) if fmt == "md" else fn(html_path)
            elif fmt in ("word", "docx"):
                path = fn(html_path, html=html)
            else:
                path = fn(html_path, html=html)
            out[fmt] = path
            log.info("[formats] %s 转换成功: %s", fmt, path)
        except Exception as e:
            log.warning("[formats] %s 转换失败: %s", fmt, e)
    return out


def cleanup_html(html_path, keep_md_imgs=True):
    """未选择 html 时: 删除 html + images/(md 选了则保留 mdimgs)"""
    art_dir = os.path.dirname(html_path)
    try:
        if os.path.isfile(html_path):
            os.remove(html_path)
        for sub in ("images", "assets"):
            dd = os.path.join(art_dir, sub)
            if os.path.isdir(dd):
                shutil.rmtree(dd, ignore_errors=True)
        # Edge 整页保存的 *_files 目录
        for _d in os.listdir(art_dir) if os.path.isdir(art_dir) else []:
            if _d.endswith("_files") and os.path.isdir(os.path.join(art_dir, _d)):
                shutil.rmtree(os.path.join(art_dir, _d), ignore_errors=True)
        log.info("[formats] 已清理中间html+images+assets: %s", art_dir)
    except Exception as e:
        log.warning("[formats] 清理失败: %s", e)